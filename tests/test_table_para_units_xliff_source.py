"""OPP#80 Wave 3: XLIFF ``<source>`` parity for the table-paragraph flag.

The Wave 2 files already pin, at the extractor level, that
``OPP_TABLE_PARAGRAPH_UNITS`` OFF yields one joined whole-cell unit and ON
yields one ``_para{p}`` unit per raw paragraph, plus the single-paragraph
(bare even when ON) and all-empty (no unit) cases. Those are intentionally not
duplicated here.

This file covers the one genuine remaining gap: the **serialized XLIFF
``<source>`` content** handed to OL/ORF. OFF must emit a single bare
``table_0_r0_c0`` whose source is the paragraph texts joined with ``"\\n"``
(the newline ORF's whole-cell path splits on); ON must emit three suffixed
units whose sources are the individual paragraphs. This is the byte-level form
that crosses the OPP→OL→ORF boundary.
"""

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from opp.extractors.docx import DOCXExtractor
from opp.extractors.pptx import PPTXExtractor
from opp.xliff import XLIFFFileGenerator

_FLAG = "OPP_TABLE_PARAGRAPH_UNITS"
_TEXTS = ("P0", "P1", "P2")


def _build_docx_cell(path: Path) -> None:
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = _TEXTS[0]
    for text in _TEXTS[1:]:
        cell.add_paragraph(text)
    doc.save(str(path))


def _build_pptx_cell(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(2))
    text_frame = frame.table.cell(0, 0).text_frame
    text_frame.text = _TEXTS[0]
    for text in _TEXTS[1:]:
        text_frame.add_paragraph().text = text
    prs.save(str(path))


def _extract(source: Path, fmt: str):
    extractor = DOCXExtractor() if fmt == "docx" else PPTXExtractor()
    return extractor.extract(source)


def _xliff_resname_sources(result, out_path: Path) -> dict[str, str]:
    XLIFFFileGenerator.from_extraction_result(result, "en", "zh").write_to_file(out_path)
    root = ET.fromstring(out_path.read_bytes())
    mapping: dict[str, str] = {}
    for unit in root.iter():
        if not unit.tag.endswith("trans-unit"):
            continue
        resname = unit.get("resname")
        if resname is None:
            continue
        source = ""
        for child in unit.iter():
            if child.tag.endswith("source"):
                source = "".join(child.itertext())
                break
        mapping[resname] = source
    return mapping


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_flag_off_vs_on_xliff_source_para_units(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fmt: str
):
    """Same file, both flag states: OFF is one ``\\n``-joined unit, ON is three."""
    source = tmp_path / f"cell.{fmt}"
    if fmt == "docx":
        _build_docx_cell(source)
    else:
        _build_pptx_cell(source)

    monkeypatch.delenv(_FLAG, raising=False)
    off = _xliff_resname_sources(
        _extract(source, fmt), tmp_path / f"off.{fmt}.xlf"
    )
    assert off == {"table_0_r0_c0": "P0\nP1\nP2"}

    monkeypatch.setenv(_FLAG, "1")
    on = _xliff_resname_sources(
        _extract(source, fmt), tmp_path / f"on.{fmt}.xlf"
    )
    assert on == {
        "table_0_r0_c0_para0": "P0",
        "table_0_r0_c0_para1": "P1",
        "table_0_r0_c0_para2": "P2",
    }
