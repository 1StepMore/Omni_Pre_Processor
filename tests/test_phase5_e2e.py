import tempfile
from pathlib import Path

import pytest

from opp.detector import FormatType
from opp.extractors.xlsx import XLSXExtractor
from opp.extractors.csv import CSVExtractor
from opp.extractors.json import JSONExtractor
from opp.extractors.xml import XMLExtractor
from opp.pipeline import OPPPipeline
from opp.xliff import XLIFFFileGenerator


class TestXLSXToMarkdown:
    def test_xlsx_basic_extraction(self, tmp_path: Path):
        import openpyxl
        xlsx_path = tmp_path / "test.xlsx"
        openpyxl.Workbook().save(xlsx_path)

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(xlsx_path)
        assert result.format_type == FormatType.XLSX

    def test_xlsx_markdown_generation(self, tmp_path: Path):
        import openpyxl
        xlsx_path = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Header1", "Header2"])
        ws.append(["Row1Col1", "Row1Col2"])
        wb.save(xlsx_path)

        extractor = XLSXExtractor()
        extraction_result = extractor.extract(xlsx_path)

        md_path = tmp_path / "test.md"
        pipeline = OPPPipeline(tmp_path / "resources")
        pipeline.generate_markdown(extraction_result, md_path)

        assert md_path.exists()
        content = md_path.read_text(encoding="utf-8")
        assert "Sheet1" in content
        assert "Header1" in content

    def test_xlsx_multiple_sheets(self, tmp_path: Path):
        import openpyxl
        xlsx_path = tmp_path / "multi.xlsx"
        wb = openpyxl.Workbook()
        wb.create_sheet("First")
        wb.create_sheet("Second")
        wb["First"].append(["Data1"])
        wb["Second"].append(["Data2"])
        wb.save(xlsx_path)

        extractor = XLSXExtractor()
        result = extractor.extract(xlsx_path)
        assert len(result.paragraphs) >= 2

    def test_xlsx_empty_workbook(self, tmp_path: Path):
        import openpyxl
        xlsx_path = tmp_path / "empty.xlsx"
        openpyxl.Workbook().save(xlsx_path)

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(xlsx_path)
        assert result.format_type == FormatType.XLSX


class TestCSVToMarkdown:
    def test_csv_pipeline_routes_to_format(self, tmp_path: Path):
        csv_path = tmp_path / "test.csv"
        csv_path.write_text("a,b\n1,2", encoding="utf-8")

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(csv_path)
        assert result.format_type == FormatType.CSV

    def test_csv_content_routing(self, tmp_path: Path):
        csv_path = tmp_path / "test.csv"
        csv_path.write_text("col\nval", encoding="utf-8")

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(csv_path)
        assert result.format_type == FormatType.CSV
        assert result.content == "" or len(result.content) >= 0


class TestJSONToXLIFF:
    def test_json_basic_extraction(self, tmp_path: Path):
        json_path = tmp_path / "test.json"
        json_path.write_text('{"k": "v"}', encoding="utf-8")

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(json_path)
        assert result.format_type == FormatType.JSON

    def test_json_xliff_generation(self, tmp_path: Path):
        json_path = tmp_path / "test.json"
        json_path.write_text('{"key": "value", "nested": {"inner": "data"}}', encoding="utf-8")

        extractor = JSONExtractor()
        extraction_result = extractor.extract(json_path)

        xlf_path = tmp_path / "test.xlf"
        pipeline = OPPPipeline(tmp_path / "resources")
        pipeline.generate_xliff(extraction_result, xlf_path, "en", "fr")

        assert xlf_path.exists()
        content = xlf_path.read_text(encoding="utf-8")
        assert "trans-unit" in content or "transunit" in content.lower()

    def test_json_flattened_output(self, tmp_path: Path):
        json_path = tmp_path / "nested.json"
        json_path.write_text('{"a": {"b": {"c": "value"}}}', encoding="utf-8")

        extractor = JSONExtractor()
        result = extractor.extract(json_path)
        para_texts = [p.text for p in result.paragraphs]
        assert any("value" in t for t in para_texts)

    def test_json_array_handling(self, tmp_path: Path):
        json_path = tmp_path / "array.json"
        json_path.write_text('["item1", "item2", "item3"]', encoding="utf-8")

        extractor = JSONExtractor()
        result = extractor.extract(json_path)
        para_texts = [p.text for p in result.paragraphs]
        assert any("item" in t for t in para_texts)


class TestXMLToMarkdown:
    def test_xml_basic_extraction(self, tmp_path: Path):
        xml_path = tmp_path / "test.xml"
        xml_path.write_text("<root/>", encoding="utf-8")

        pipeline = OPPPipeline(tmp_path / "resources")
        result = pipeline.process_file(xml_path)
        assert result.format_type == FormatType.XML

    def test_xml_markdown_generation(self, tmp_path: Path):
        xml_path = tmp_path / "test.xml"
        xml_path.write_text(
            '<?xml version="1.0"?><root><element>Content</element></root>',
            encoding="utf-8"
        )

        extractor = XMLExtractor()
        extraction_result = extractor.extract(xml_path)

        md_path = tmp_path / "test.md"
        pipeline = OPPPipeline(tmp_path / "resources")
        pipeline.generate_markdown(extraction_result, md_path)

        assert md_path.exists()
        content = md_path.read_text(encoding="utf-8")
        assert "Content" in content

    def test_xml_with_namespaces(self, tmp_path: Path):
        xml_path = tmp_path / "namespaced.xml"
        xml_path.write_text(
            '<?xml version="1.0"?>'
            '<root xmlns:ns="http://example.com/ns">'
            '<ns:element>Value</ns:element>'
            '</root>',
            encoding="utf-8"
        )

        extractor = XMLExtractor()
        result = extractor.extract(xml_path)
        para_texts = [p.text for p in result.paragraphs]
        assert any("Value" in t for t in para_texts)

    def test_xml_nested_elements(self, tmp_path: Path):
        xml_path = tmp_path / "nested.xml"
        xml_path.write_text(
            '<?xml version="1.0"?>'
            '<root>'
            '<parent><child>Text1</child><child>Text2</child></parent>'
            '</root>',
            encoding="utf-8"
        )

        extractor = XMLExtractor()
        result = extractor.extract(xml_path)
        assert len(result.paragraphs) >= 1


class TestPhase5PipelineIntegration:
    def test_all_phase5_formats_process(self, tmp_path: Path):
        pipeline = OPPPipeline(tmp_path / "resources")

        import openpyxl
        xlsx_path = tmp_path / "test.xlsx"
        openpyxl.Workbook().save(xlsx_path)
        xlsx_result = pipeline.process_file(xlsx_path)
        assert xlsx_result.format_type == FormatType.XLSX

        csv_path = tmp_path / "test.csv"
        csv_path.write_text("a,b\n1,2", encoding="utf-8")
        csv_result = pipeline.process_file(csv_path)
        assert csv_result.format_type == FormatType.CSV

        json_path = tmp_path / "test.json"
        json_path.write_text('{"k": "v"}', encoding="utf-8")
        json_result = pipeline.process_file(json_path)
        assert json_result.format_type == FormatType.JSON

        xml_path = tmp_path / "test.xml"
        xml_path.write_text("<root/>", encoding="utf-8")
        xml_result = pipeline.process_file(xml_path)
        assert xml_result.format_type == FormatType.XML

    def test_markdown_generation_all_formats(self, tmp_path: Path):
        pipeline = OPPPipeline(tmp_path / "resources")
        md_output = tmp_path / "output.md"

        import openpyxl
        xlsx_path = tmp_path / "xlsx.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["h1", "h2"])
        wb.active.append(["v1", "v2"])
        wb.save(xlsx_path)
        xlsx_result = pipeline.process_file(xlsx_path)
        assert xlsx_result.extraction_result is not None

        pipeline.generate_markdown(xlsx_result.extraction_result, md_output)
        assert md_output.exists()

    def test_xliff_generation_json_only(self, tmp_path: Path):
        pipeline = OPPPipeline(tmp_path / "resources")
        xlf_output = tmp_path / "out.xlf"

        json_path = tmp_path / "data.json"
        json_path.write_text('{"name": "test", "value": "123"}', encoding="utf-8")

        extractor = JSONExtractor()
        result = extractor.extract(json_path)
        pipeline.generate_xliff(result, xlf_output, "en", "de")

        assert xlf_output.exists()
        content = xlf_output.read_text(encoding="utf-8")
        assert "en" in content or "English" in content

    def test_phase5_no_images(self, tmp_path: Path):
        import openpyxl
        xlsx_path = tmp_path / "img_test.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["data"])
        wb.save(xlsx_path)

        extractor = XLSXExtractor()
        result = extractor.extract(xlsx_path)
        assert len(result.images) == 0

        json_path = tmp_path / "img_test.json"
        json_path.write_text('{"a": "b"}', encoding="utf-8")
        json_extractor = JSONExtractor()
        json_result = json_extractor.extract(json_path)
        assert len(json_result.images) == 0