"""OPP#80 Wave 2: DOCX per-paragraph table-cell units.

``OPP_TABLE_PARAGRAPH_UNITS`` gates the cell granularity. OFF (default) keeps
the legacy one-unit-per-cell behavior byte-for-byte. ON emits one
``TableCellData`` per NON-EMPTY paragraph of any cell with **>= 2 raw
direct-child ``w:p``**, each carrying ``para_index`` — the 0-based index into
the RAW ``w:p`` list, empty paragraphs included, exactly the list ORF resolves
with ``tc.findall(w:p)``. An empty paragraph consumes an index but emits no
unit, so ``[P0, "", P2]`` yields indices 0 and 2. A single-paragraph cell
keeps the bare whole-cell unit.

These tests drive the real ``DOCXExtractor`` and ``XLIFFFileGenerator`` and
patch the flag through its real seam (the env var, read by the helper at call
time).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from opp.extractors.docx import DOCXExtractor
from opp.xliff import XLIFFFileGenerator


def _cell_docx(tmp_path: Path, texts: list[str]) -> Path:
    """A one-cell DOCX whose only cell has one raw ``w:p`` per entry in texts."""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = texts[0]
    for text in texts[1:]:
        cell.add_paragraph(text)
    path = tmp_path / "cell.docx"
    doc.save(str(path))
    return path


def _extract_cells(tmp_path: Path, texts: list[str]):
    return DOCXExtractor().extract(_cell_docx(tmp_path, texts)).table_cells


def _generate_xliff(result, tmp_path: Path) -> bytes:
    out_path = tmp_path / "out.xlf"
    gen = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
    gen.write_to_file(out_path)
    return out_path.read_bytes()


def _resnames(xliff_bytes: bytes) -> list[str]:
    return [r.decode("utf-8") for r in re.findall(rb'resname="([^"]+)"', xliff_bytes)]


class TestDOCXTableParagraphUnits:
    def test_flag_off_three_paragraph_cell_is_one_joined_unit(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag OFF pins legacy: ONE unit, paragraphs joined by ``\\n``."""
        monkeypatch.delenv("OPP_TABLE_PARAGRAPH_UNITS", raising=False)
        cells = _extract_cells(tmp_path, ["P0", "P1", "P2"])

        assert len(cells) == 1
        assert cells[0].text == "P0\nP1\nP2"
        assert cells[0].para_index is None

    def test_flag_on_three_paragraph_cell_yields_three_indexed_units(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON splits a 3-paragraph cell into para_index 0, 1, 2."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        cells = _extract_cells(tmp_path, ["P0", "P1", "P2"])

        assert [c.para_index for c in cells] == [0, 1, 2]
        assert [c.text for c in cells] == ["P0", "P1", "P2"]

    def test_flag_on_single_paragraph_cell_stays_bare(
        self, tmp_path: Path, monkeypatch
    ):
        """A one-paragraph cell keeps the bare whole-cell unit (min churn)."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        cells = _extract_cells(tmp_path, ["Only"])

        assert len(cells) == 1
        assert cells[0].text == "Only"
        assert cells[0].para_index is None

    def test_flag_on_empty_paragraph_consumes_but_emits_nothing(
        self, tmp_path: Path, monkeypatch
    ):
        """ALIGNMENT: ``[P0, "", P2]`` -> indices 0 and 2, never 0 and 1."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        cells = _extract_cells(tmp_path, ["P0", "", "P2"])

        assert [c.para_index for c in cells] == [0, 2]
        assert [c.text for c in cells] == ["P0", "P2"]

    def test_flag_on_all_empty_cell_emits_no_unit(
        self, tmp_path: Path, monkeypatch
    ):
        """A cell whose paragraphs are ALL empty produces no unit."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        cells = _extract_cells(tmp_path, ["", "", ""])

        assert cells == []

    def test_generator_resnames_switch_with_flag(
        self, tmp_path: Path, monkeypatch
    ):
        """End-to-end: the generator appends ``_para{p}`` only when set."""
        docx_path = _cell_docx(tmp_path, ["P0", "P1", "P2"])

        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        on_result = DOCXExtractor().extract(docx_path)
        on_resnames = [n for n in _resnames(_generate_xliff(on_result, tmp_path))
                       if n.startswith("table_")]
        assert on_resnames == [
            "table_0_r0_c0_para0",
            "table_0_r0_c0_para1",
            "table_0_r0_c0_para2",
        ]

        monkeypatch.delenv("OPP_TABLE_PARAGRAPH_UNITS", raising=False)
        off_result = DOCXExtractor().extract(docx_path)
        off_resnames = [n for n in _resnames(_generate_xliff(off_result, tmp_path))
                        if n.startswith("table_")]
        assert off_resnames == ["table_0_r0_c0"]
