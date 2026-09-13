"""T-04: OPP CLI machine-readable ``--json`` output.

Contract (gap register T-04, ``.omo/plans/agent-oriented-gap-register.md``):

  - ``opp <fixture> --json`` writes **exactly one** parseable JSON object to
    stdout (human logs/progress never pollute stdout).
  - The object carries ``success``, the emitted output paths
    (``md`` / ``xliff`` / ``skeleton`` / ``manifest`` as applicable), and
    ``warnings``.
  - Every reported path must exist on disk (no misleading success output).
  - A missing or malformed input returns a JSON error object **and** a
    non-zero exit code.
  - Plain text mode (no ``--json``) is unchanged.

These tests spawn the real CLI as a subprocess (``python -m opp.cli``), the
same surface an agent hits from the shell.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest


OPP_DIR = Path(__file__).resolve().parents[1]


def _make_docx(path: Path, text: str = "T-04 fixture paragraph about extraction.") -> Path:
    """Build a minimal but valid DOCX fixture."""
    from docx import Document

    doc = Document()
    doc.add_heading("T-04 Heading", level=1)
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def run_opp(*args: str, cwd: Path | None = None) -> subprocess.CompletedProcess:
    """Invoke the OPP CLI in a subprocess with a real stdout/stderr split."""
    return subprocess.run(
        [sys.executable, "-m", "opp.cli", *args],
        capture_output=True,
        text=True,
        cwd=str(cwd or OPP_DIR),
        timeout=180,
    )


class TestSingleFileJson:
    """``opp <fixture> --json`` → one JSON object with output paths."""

    def test_json_is_exactly_one_object_with_outputs_and_warnings(self, tmp_path):
        # Given: a valid DOCX and an output dir
        src = _make_docx(tmp_path / "sample.docx")
        out = tmp_path / "out"

        # When: the CLI runs with --json (md path)
        proc = run_opp(
            str(src), "--target-format", "md",
            "--source-lang", "en", "--target-lang", "zh",
            "--output-dir", str(out), "--json",
        )

        # Then: stdout is exactly one parseable JSON object
        assert proc.returncode == 0, (
            f"exit={proc.returncode}\nstderr={proc.stderr}\nstdout={proc.stdout}"
        )
        payload = json.loads(proc.stdout)  # raises if >1 object or trailing text
        assert payload["success"] is True
        assert "warnings" in payload

        # And: the reported paths exist (no misleading success output)
        outputs = payload["outputs"]
        assert {"md", "manifest"} <= set(outputs), outputs
        for key in ("md", "manifest"):
            assert Path(outputs[key]).exists(), f"reported {key} missing: {outputs[key]}"

    def test_json_both_reports_md_xliff_skeleton_manifest(self, tmp_path):
        # Given: a DOCX (skeletonizable) and target-format both
        src = _make_docx(tmp_path / "sample.docx")
        out = tmp_path / "out"

        # When
        proc = run_opp(
            str(src), "--target-format", "both",
            "--source-lang", "en", "--target-lang", "zh",
            "--output-dir", str(out), "--json",
        )

        # Then: all four output classes are reported and exist
        assert proc.returncode == 0, f"stderr={proc.stderr}\nstdout={proc.stdout}"
        payload = json.loads(proc.stdout)
        assert payload["success"] is True
        outputs = payload["outputs"]
        assert {"md", "xliff", "skeleton", "manifest"} <= set(outputs), outputs
        for key in ("md", "xliff", "skeleton", "manifest"):
            assert Path(outputs[key]).exists(), f"reported {key} missing: {outputs[key]}"


class TestBatchJson:
    """``opp --batch <a> <b> --json`` → one JSON object covering every file."""

    def test_batch_emits_one_object_with_all_files(self, tmp_path):
        # Given: two DOCX inputs
        a = _make_docx(tmp_path / "a.docx", "First document.")
        b = _make_docx(tmp_path / "b.docx", "Second document.")
        out = tmp_path / "out"

        # When
        proc = run_opp(
            str(a), str(b), "--batch", "--target-format", "md",
            "--source-lang", "en", "--target-lang", "zh",
            "--output-dir", str(out), "--json",
        )

        # Then: exactly one object, both files present, every path exists
        assert proc.returncode == 0, f"stderr={proc.stderr}\nstdout={proc.stdout}"
        payload = json.loads(proc.stdout)
        assert payload["success"] is True
        assert len(payload["files"]) == 2
        for entry in payload["files"]:
            assert entry["success"] is True
            assert entry["outputs"], entry
            for key, p in entry["outputs"].items():
                assert Path(p).exists(), f"{entry['file']}: reported {key} missing: {p}"


class TestJsonFailurePaths:
    """Failure must be a JSON error object with a non-zero exit code."""

    def test_missing_file_json_error_nonzero(self, tmp_path):
        # Given: a path that does not exist
        missing = tmp_path / "does_not_exist.docx"

        # When
        proc = run_opp(
            str(missing), "--target-format", "md",
            "--output-dir", str(tmp_path / "out"), "--json",
        )

        # Then: non-zero exit + parseable JSON error on stdout
        assert proc.returncode != 0, f"stdout={proc.stdout}\nstderr={proc.stderr}"
        payload = json.loads(proc.stdout)
        assert payload["success"] is False
        assert payload.get("error"), payload

    def test_corrupt_docx_json_error_nonzero(self, tmp_path):
        # Given: a file with a .docx extension but garbage content
        bad = tmp_path / "corrupt.docx"
        bad.write_bytes(b"not a valid docx at all")

        # When
        proc = run_opp(
            str(bad), "--target-format", "md",
            "--output-dir", str(tmp_path / "out"), "--json",
        )

        # Then
        assert proc.returncode != 0, f"stdout={proc.stdout}\nstderr={proc.stderr}"
        payload = json.loads(proc.stdout)
        assert payload["success"] is False

    def test_missing_target_lang_json_error_nonzero(self, tmp_path):
        # Given: xlf requested without --target-lang
        src = _make_docx(tmp_path / "sample.docx")

        # When
        proc = run_opp(
            str(src), "--target-format", "xlf", "--json",
            "--output-dir", str(tmp_path / "out"),
        )

        # Then: structured JSON error, not argparse usage text on stdout
        assert proc.returncode != 0
        payload = json.loads(proc.stdout)
        assert payload["success"] is False
        assert "target-lang" in json.dumps(payload).lower()


class TestStaleState:
    """A previous run's artifact must not be reported as this run's output."""

    def test_json_reports_fresh_output_after_stale_file(self, tmp_path):
        # Given: a stale .md pre-planted in the output dir
        src = _make_docx(tmp_path / "sample.docx")
        out = tmp_path / "out"
        out.mkdir()
        (out / "sample.md").write_text("STALE CONTENT", encoding="utf-8")

        # When
        proc = run_opp(
            str(src), "--target-format", "md",
            "--source-lang", "en", "--target-lang", "zh",
            "--output-dir", str(out), "--json",
        )

        # Then: the reported md path exists and was regenerated
        assert proc.returncode == 0, f"stderr={proc.stderr}"
        payload = json.loads(proc.stdout)
        md = Path(payload["outputs"]["md"])
        assert md.exists()
        assert "STALE CONTENT" not in md.read_text(encoding="utf-8")


class TestTextModeUnchanged:
    """Characterization: without --json the CLI keeps its current contract."""

    def test_text_mode_stdout_is_not_json(self, tmp_path):
        # Given
        src = _make_docx(tmp_path / "sample.docx")
        out = tmp_path / "out"

        # When: no --json
        proc = run_opp(
            str(src), "--target-format", "md",
            "--source-lang", "en", "--target-lang", "zh",
            "--output-dir", str(out),
        )

        # Then: exit 0, artifacts written, stdout is not a JSON document
        assert proc.returncode == 0, f"stderr={proc.stderr}"
        assert (out / "sample.md").exists()
        with pytest.raises(json.JSONDecodeError):
            json.loads(proc.stdout)


class TestJsonHelp:
    def test_help_documents_json_flag(self):
        proc = run_opp("--help")
        assert proc.returncode == 0
        assert "--json" in proc.stdout
