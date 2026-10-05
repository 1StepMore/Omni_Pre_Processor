"""XLSX skeletons carry an ``xliff_map.json`` sidecar (1StepMore/Omni_Pre_Processor#92).

The skeleton is the ORIGINAL workbook bytes plus one added archive entry, so
Excel safety holds without touching a single ``xl/*.xml`` byte. The map tells
ORF which spreadsheet cell each XLIFF trans-unit row came from and which of
those cells may be written back.

The matrix fixture (a single ``A1="Hello"``, one sheet, no formulas, no merges)
cannot exercise any of that, so the fixture here is deliberately rich: two
sheets, a formula cell carrying a cached value, a merged range, a numeric cell,
a date, a blank cell, an all-blank row and a multi-column row.
"""

import json
import re
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from lxml import etree
from opp.extractors.xlsx import XLSXExtractor
from opp.pipeline import OPPPipeline
from opp.xliff.generator import XLIFFFileGenerator

XLIFF_NS = {"x": "urn:oasis:names:tc:xliff:document:1.2"}
SHEET1_XML = "xl/worksheets/sheet1.xml"
# Spelled out rather than imported so this module still imports against a source
# tree that has no skeleton at all — the tests must fail on behaviour, not on a
# missing symbol.
XLIFF_MAP_ENTRY = "xliff_map.json"


def _patch_sheet_xml(path: Path, patch) -> None:
    with zipfile.ZipFile(path) as zf:
        entries = [(i, zf.read(i)) for i in zf.infolist()]
    out = []
    for info, data in entries:
        if info.filename == SHEET1_XML:
            data = patch(data.decode("utf-8")).encode("utf-8")
        out.append((info, data))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for info, data in out:
            zf.writestr(info, data)


def _inject_cached_value(path: Path, cell_ref: str, cached: str) -> None:
    """Store the cached value Excel keeps for a formula cell.

    openpyxl writes ``<c r="B2"><f>SUM(1,1)</f><v></v></c>`` — an empty ``<v>``,
    so ``load_workbook(data_only=True)`` reports ``None``. No Excel-authored
    workbook looks like that, and in that state the formula cell's extracted
    "text" cannot be observed at all, which is exactly the state the map's
    ``translatable`` rule exists to guard.
    """

    def patch(xml: str) -> str:
        xml, count = re.subn(
            rf'(<c r="{cell_ref}"[^>]*>.*?</f>)<v></v>',
            lambda m: f"{m.group(1)}<v>{cached}</v>",
            xml,
            count=1,
        )
        assert count == 1, f"no formula cell {cell_ref} in {SHEET1_XML}"
        return xml

    _patch_sheet_xml(path, patch)


def _inject_error_cell(path: Path, row: int, cell_ref: str, literal: str) -> None:
    """Add a real ``t="e"`` error cell, which openpyxl cannot author."""

    def patch(xml: str) -> str:
        cell = f'<c r="{cell_ref}" t="e"><v>{literal}</v></c>'
        xml, count = re.subn(
            rf'(<row r="{row}"[^>]*>)', lambda m: m.group(1) + cell, xml, count=1
        )
        assert count == 1, f"no row {row} in {SHEET1_XML}"
        return xml

    _patch_sheet_xml(path, patch)


@pytest.fixture
def rich_xlsx(tmp_path: Path) -> Path:
    """Two sheets exercising every ``translatable=false`` rule."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    # Row 1: multi-column row — one translatable string plus a numeric cell.
    ws.append(["Product Name", 42])
    # Row 2: formula cell; its extracted text is the cached value.
    ws["A2"] = "Computed Label"
    ws["B2"] = "=SUM(1,1)"
    # Row 3: merged range followed by more text, so the non-anchor cells stay
    # inside the row instead of being stripped off its right edge.
    ws.merge_cells("B3:D3")
    ws["A3"] = "Merged Anchor"
    ws["B3"] = "Merged Block"
    ws["E3"] = "Row Three Tail"
    # Row 4: date cell, blank cell, then text.
    ws["A4"] = date(2024, 1, 15)
    ws["C4"] = "Date Row Tail"
    # Row 5: entirely blank row.
    ws.append([None, None, None, None, None])
    # Row 6: single text cell.
    ws["A6"] = "Total Amount"

    ws2 = wb.create_sheet("Notes")
    ws2.append(["Second Sheet Text"])

    path = tmp_path / "rich.xlsx"
    wb.save(str(path))
    _inject_cached_value(path, "B2", "2")
    _inject_error_cell(path, 4, "D4", "#DIV/0!")
    return path


def _write_skeleton(result, tmp_path: Path) -> Path:
    output_dir = tmp_path / "output"
    output_dir.mkdir(exist_ok=True)
    skeleton_path = OPPPipeline(
        resource_storage_dir=tmp_path / "resources"
    ).save_skeleton(result, "rich", output_dir)
    assert skeleton_path is not None, "XLSX must now produce a skeleton.zip"
    assert skeleton_path.name == "rich.skeleton.zip"
    return skeleton_path


def _map_from_skeleton(result, tmp_path: Path) -> dict:
    skeleton_path = _write_skeleton(result, tmp_path)
    with zipfile.ZipFile(skeleton_path) as zf:
        assert XLIFF_MAP_ENTRY in zf.namelist(), "sidecar missing from skeleton"
        return json.loads(zf.read(XLIFF_MAP_ENTRY).decode("utf-8"))


def _cell_index(xliff_map: dict, sheet: str) -> dict[str, bool]:
    """Every cell of a sheet keyed by its spreadsheet ref."""
    return {
        cell["ref"]: cell["translatable"]
        for row in xliff_map["sheets"][sheet]
        for cell in row["cells"]
    }


class TestXLSXSkeletonZip:
    def test_skeleton_keeps_every_original_entry_byte_identical(
        self, rich_xlsx, tmp_path
    ):
        """No ``xl/*.xml`` byte is rewritten — that is what makes Excel safe."""
        result = XLSXExtractor().extract(rich_xlsx)
        skeleton_path = _write_skeleton(result, tmp_path)

        with zipfile.ZipFile(rich_xlsx) as original:
            original_names = original.namelist()
            original_bytes = {n: original.read(n) for n in original_names}
        with zipfile.ZipFile(skeleton_path) as skeleton:
            assert skeleton.testzip() is None
            skeleton_names = skeleton.namelist()
            skeleton_bytes = {n: skeleton.read(n) for n in skeleton_names}

        for name in original_names:
            assert name in skeleton_names, f"{name} dropped from skeleton"
            assert skeleton_bytes[name] == original_bytes[name], (
                f"{name} is not byte-identical to the original workbook entry"
            )

        assert set(skeleton_names) - set(original_names) == {XLIFF_MAP_ENTRY}
        assert "xl/workbook.xml" in skeleton_names, (
            "ORF's FormatDetector.detect_from_skeleton keys XLSX off xl/workbook.xml"
        )

    def test_map_schema_version_and_sheets(self, rich_xlsx, tmp_path):
        result = XLSXExtractor().extract(rich_xlsx)
        xliff_map = _map_from_skeleton(result, tmp_path)

        assert xliff_map["version"] == 1
        assert set(xliff_map["sheets"]) == {"Data", "Notes"}
        for sheet_rows in xliff_map["sheets"].values():
            for row in sheet_rows:
                assert set(row) == {"id", "row", "cells"}
                assert set(row["cells"][0]) == {"ref", "translatable"}

    def test_translatable_flags_per_cell_rule(self, rich_xlsx, tmp_path):
        result = XLSXExtractor().extract(rich_xlsx)
        flags = _cell_index(_map_from_skeleton(result, tmp_path), "Data")

        # Plain text is translatable.
        assert flags["A1"] is True
        assert flags["A6"] is True
        # Numeric cell of the multi-column row: data, not prose.
        assert flags["B1"] is False
        # Formula cell: writing a translation back would destroy the formula.
        assert flags["B2"] is False
        # Merged range: only the top-left anchor holds the value.
        assert flags["B3"] is True
        assert flags["C3"] is False
        assert flags["D3"] is False
        # Text outside the merged range in the same row stays translatable.
        assert flags["A3"] is True
        assert flags["E3"] is True
        # Date cell and blank cell are not translatable text.
        assert flags["A4"] is False
        assert flags["B4"] is False
        assert flags["C4"] is True
        # Excel error literal: reads as a string, is not text.
        assert flags["D4"] is False
        assert flags["E4"] is False

    def test_formula_inventory_failure_disables_the_map(self, rich_xlsx, monkeypatch):
        """Fail closed: an untrustworthy formula inventory means no skeleton."""
        real_load = openpyxl.load_workbook

        def flaky(*args, **kwargs):
            if kwargs.get("read_only"):
                raise OSError("cannot read")
            return real_load(*args, **kwargs)

        monkeypatch.setattr(openpyxl, "load_workbook", flaky)
        result = XLSXExtractor().extract(rich_xlsx)
        assert result.skeleton is None
        assert result.skeleton_files is None

    def test_row_numbers_and_per_sheet_ids(self, rich_xlsx, tmp_path):
        result = XLSXExtractor().extract(rich_xlsx)
        xliff_map = _map_from_skeleton(result, tmp_path)

        assert [row["row"] for row in xliff_map["sheets"]["Notes"]] == [1]
        assert [row["row"] for row in xliff_map["sheets"]["Data"]] == [1, 2, 3, 4, 5, 6]
        ids = [row["id"] for rows in xliff_map["sheets"].values() for row in rows]
        assert len(ids) == len(set(ids)), "ids must be unique across sheets"

    def test_blank_row_maps_to_all_false_cells(self, rich_xlsx, tmp_path):
        """An all-blank row still yields a trans-unit (``" |  |"`` survives the
        strip), so the map must list it with nothing writable."""
        result = XLSXExtractor().extract(rich_xlsx)
        xliff_map = _map_from_skeleton(result, tmp_path)

        blank = next(r for r in xliff_map["sheets"]["Data"] if r["row"] == 5)
        assert blank["cells"] == [
            {"ref": f"{col}5", "translatable": False} for col in "ABCDE"
        ]

    def test_map_ids_match_generated_xliff_trans_units(self, rich_xlsx, tmp_path):
        """The ids come off the shared paragraph list, so each one must name the
        trans-unit the generator really emits, with the very same source text."""
        result = XLSXExtractor().extract(rich_xlsx)
        xliff_map = _map_from_skeleton(result, tmp_path)

        xliff_path = tmp_path / "rich.xlf"
        XLIFFFileGenerator.from_extraction_result(result, "en", "zh").write_to_file(
            xliff_path
        )
        tree = etree.parse(str(xliff_path))
        units = {
            tu.get("id"): tu.findtext("x:source", "", namespaces=XLIFF_NS)
            for tu in tree.xpath("//x:trans-unit", namespaces=XLIFF_NS)
        }

        assert len(units) == len(result.paragraphs)
        for sheet_rows in xliff_map["sheets"].values():
            for row in sheet_rows:
                unit_id = str(row["id"])
                assert unit_id in units, f"map id {unit_id} has no trans-unit"
                assert units[unit_id] == result.paragraphs[row["id"] - 1].text

    def test_md_output_row_granularity_unchanged(self, rich_xlsx, tmp_path):
        """Row granularity is the MD contract; the map must not reshape it."""
        result = XLSXExtractor().extract(rich_xlsx)
        texts = [p.text for p in result.paragraphs]

        assert texts[0] == "=== Sheet: Data ==="
        assert texts[1] == "Product Name | 42 |  |  |"
        assert texts[2] == "Computed Label | 2 |  |  |"
        assert texts[3] == "Merged Anchor | Merged Block |  |  | Row Three Tail"
        assert texts[4] == "2024-01-15T00:00:00 |  | Date Row Tail | #DIV/0! |"
        assert texts[5] == "|  |  |  |"
        assert texts[6] == "Total Amount |  |  |  |"
        assert texts[7] == "=== Sheet: Notes ==="
        assert texts[8] == "Second Sheet Text"

    def test_manifest_key_files_describe_the_workbook(self, rich_xlsx, tmp_path):
        result = XLSXExtractor().extract(rich_xlsx)
        assert result.skeleton_files is not None
        assert "xl/workbook.xml" in result.skeleton_files
        assert "[Content_Types].xml" in result.skeleton_files
        assert "xl/worksheets/sheet1.xml" in result.skeleton_files