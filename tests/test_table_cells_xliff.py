"""Issue A regression: table cells must become XLIFF trans-units.

DOCX ``w:tc`` and HTML ``<td>``/``<th>`` content used to be dropped from the
XLIFF entirely, so ``apply_xliff`` left the whole table untranslated. These
tests drive the real extractors and the real ``XLIFFFileGenerator`` and assert
on the emitted XLIFF bytes / resnames (``table_{t}_r{r}_c{c}``).
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from pathlib import Path

from opp.extractors.docx import DOCXExtractor
from opp.extractors.html import HTMLExtractor
from opp.xliff import XLIFFFileGenerator


def _generate_xliff(result, tmp_path: Path) -> bytes:
    out_path = tmp_path / "out.xlf"
    gen = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
    gen.write_to_file(out_path)
    return out_path.read_bytes()


def _resnames(xliff_bytes: bytes) -> list[str]:
    return [r.decode("utf-8") for r in re.findall(rb'resname="([^"]+)"', xliff_bytes)]


def _unit_count(xliff_bytes: bytes) -> int:
    return xliff_bytes.count(b"<trans-unit")


def _sources(xliff_bytes: bytes) -> list[str]:
    root = ET.fromstring(xliff_bytes)
    sources = []
    for elem in root.iter():
        if elem.tag.endswith("source"):
            sources.append("".join(elem.itertext()))
    return sources


def _resname_sources(xliff_bytes: bytes) -> dict[str, str]:
    root = ET.fromstring(xliff_bytes)
    mapping: dict[str, str] = {}
    for elem in root.iter():
        if not elem.tag.endswith("trans-unit"):
            continue
        resname = elem.get("resname")
        source = ""
        for child in elem.iter():
            if child.tag.endswith("source"):
                source = "".join(child.itertext())
                break
        if resname is not None:
            mapping[resname] = source
    return mapping


class TestDOCXTableCells:
    def test_docx_table_cells_become_trans_units(self, tmp_path: Path):
        """A 2x2 table plus 2 body paragraphs yields 6 trans-units."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("First body paragraph")
        doc.add_paragraph("Second body paragraph")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        table.rows[1].cells[0].text = "C"
        table.rows[1].cells[1].text = "D"
        docx_path = tmp_path / "cells.docx"
        doc.save(str(docx_path))

        result = DOCXExtractor().extract(docx_path)
        xliff_bytes = _generate_xliff(result, tmp_path)

        assert _unit_count(xliff_bytes) == 6
        resnames = set(_resnames(xliff_bytes))
        for expected in (
            "table_0_r0_c0",
            "table_0_r0_c1",
            "table_0_r1_c0",
            "table_0_r1_c1",
        ):
            assert expected in resnames, f"missing resname {expected!r}: {resnames}"

    def test_docx_empty_cell_produces_no_unit(self, tmp_path: Path):
        """An empty cell is skipped; only the non-empty cell becomes a unit."""
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "Filled"
        docx_path = tmp_path / "empty_cell.docx"
        doc.save(str(docx_path))

        result = DOCXExtractor().extract(docx_path)
        xliff_bytes = _generate_xliff(result, tmp_path)

        assert _unit_count(xliff_bytes) == 1
        resnames = set(_resnames(xliff_bytes))
        assert "table_0_r0_c0" in resnames
        assert "table_0_r1_c0" not in resnames


class TestHTMLTableCells:
    FIXTURE = (
        "<h1>Title</h1>"
        "<table><tr><th>A</th><th>B</th></tr>"
        "<tr><td>C</td><td>D</td></tr></table>"
    )

    def test_html_table_cells_become_trans_units(self, tmp_path: Path):
        """The heading plus four cells yields 5 units; no flattened ``|`` row."""
        html_path = tmp_path / "table.html"
        html_path.write_text(self.FIXTURE, encoding="utf-8")

        result = HTMLExtractor().extract(html_path)
        xliff_bytes = _generate_xliff(result, tmp_path)

        assert _unit_count(xliff_bytes) == 5
        resnames = set(_resnames(xliff_bytes))
        for expected in (
            "table_0_r0_c0",
            "table_0_r0_c1",
            "table_0_r1_c0",
            "table_0_r1_c1",
        ):
            assert expected in resnames, f"missing resname {expected!r}: {resnames}"

        for source in _sources(xliff_bytes):
            assert not source.strip().startswith("|"), (
                f"flattened markdown table row leaked into XLIFF: {source!r}"
            )


def test_non_body_resname_keeps_paragraph_index_when_table_rows_are_skipped(tmp_path: Path):
    # HTML with a markdown-visible table: the paragraph that follows the table must keep
    # its REAL index in ExtractionResult.paragraphs in its non_body_N resname.
    html = ("<h1>T</h1><p>Before table.</p>"
            "<table><tr><th>A</th><th>B</th></tr><tr><td>C</td><td>D</td></tr></table>"
            "<p>After table.</p>")
    html_path = tmp_path / "table_index.html"
    html_path.write_text(html, encoding="utf-8")

    result = HTMLExtractor().extract(html_path)
    xliff_bytes = _generate_xliff(result, tmp_path)

    # Every paragraph that is emitted (i.e. not a flattened markdown table row)
    # must map non_body_<real paragraph index> -> its own text.
    expected = {
        f"non_body_{i}": p.text
        for i, p in enumerate(result.paragraphs)
        if i not in result.table_row_paragraph_indices
    }
    actual = {
        resname: source
        for resname, source in _resname_sources(xliff_bytes).items()
        if resname.startswith("non_body_")
    }
    assert actual == expected

    # The paragraph after the table must NOT be renumbered to follow the
    # filtered counter; its resname is derived from its real position.
    after_index = next(
        i for i, p in enumerate(result.paragraphs) if p.text == "After table."
    )
    assert f"non_body_{after_index}" in actual
    assert actual[f"non_body_{after_index}"] == "After table."

def test_pipe_prefixed_non_table_paragraph_survives_in_xliff(tmp_path: Path):
    # Regression: suppressing the flattened markdown rows of a real table must be
    # driven by what the extractor recorded, not by the paragraph starting with
    # ``|``. A text check also dropped genuine ``|``-prefixed paragraphs (code
    # blocks, ASCII art), which then never reached the translator and stayed in
    # the source language after backfill.
    #
    # A 3-column table is markdownified into ``|`` rows (so the suppression path
    # is genuinely exercised here); the ``<pre>`` line is not a table row.
    html = (
        "<p>Intro paragraph.</p>"
        "<table><tr><th>Model</th><th>Ingress</th><th>Weight</th></tr>"
        "<tr><td>X1</td><td>IP67</td><td>0.5 ms</td></tr>"
        "<tr><td>X2</td><td>IP67</td><td>1.2 ms</td></tr></table>"
        "<pre>| piped | literal line that is NOT a table |</pre>"
        "<p>Outro paragraph.</p>"
    )
    html_path = tmp_path / "pipe_paragraph.html"
    html_path.write_text(html, encoding="utf-8")

    result = HTMLExtractor().extract(html_path)
    assert result.table_row_paragraph_indices, (
        "fixture no longer produces flattened markdown table rows - the "
        "suppression path would be untested by this case"
    )

    xliff_bytes = _generate_xliff(result, tmp_path)
    sources = _sources(xliff_bytes)

    assert any("piped" in s for s in sources), (
        f"non-table pipe-prefixed paragraph was dropped from the XLIFF: {sources}"
    )

    # The suppression still works: no markdown row rendering leaked through.
    for source in sources:
        assert not source.strip().startswith("| --- |"), (
            f"flattened markdown delimiter row leaked into XLIFF: {source!r}"
        )
    for row in ("Model | Ingress | Weight", "X1 | IP67 | 0.5 ms"):
        assert not any(row in s for s in sources), (
            f"flattened markdown table row leaked into XLIFF: {row!r}"
        )

    # And every real cell is still emitted positionally.
    resnames = _resnames(xliff_bytes)
    for expected in (
        "table_0_r0_c0", "table_0_r0_c1", "table_0_r0_c2",
        "table_0_r1_c0", "table_0_r1_c1", "table_0_r1_c2",
        "table_0_r2_c0", "table_0_r2_c1", "table_0_r2_c2",
    ):
        assert expected in resnames, f"missing resname {expected!r}: {resnames}"
