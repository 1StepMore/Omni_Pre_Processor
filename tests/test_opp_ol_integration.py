"""Integration tests for OPP→OL pipeline."""

import pytest
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document

from opp.detector import FormatType
from opp.pipeline import OPPPipeline
from opp.utils.dataclasses import (
    ExtractionResult,
    ParagraphData,
    TableData,
    ImageData,
    DocumentMetadata,
)

from tests.mock_ol import MockOLTool

XLIFF_NS_1_2 = "urn:oasis:names:tc:xliff:document:1.2"
XLIFF_NS_1_1 = "urn:oasis:names:tc:xliff:document:1.1"


def detect_xliff_ns(root) -> str:
    for elem in root.iter():
        tag = elem.tag
        if tag.startswith("{"):
            ns_candidate = tag[1:tag.index("}")]
            if "xliff" in ns_candidate:
                return ns_candidate
    for uri in root.attrib.values():
        if "xliff" in uri:
            return uri
    return XLIFF_NS_1_2


def find_elem(root, ns_uri, tag):
    try:
        return next(root.iter(f"{{{ns_uri}}}{tag}"))
    except StopIteration:
        return None


class TestOPPtoOLIntegration:
    @pytest.fixture
    def temp_output_dir(self, tmp_path: Path) -> Path:
        output_dir = tmp_path / "output"
        output_dir.mkdir()
        return output_dir

    @pytest.fixture
    def sample_docx(self, tmp_path: Path) -> Path:
        doc = Document()
        doc.add_heading("Test Document", level=1)
        doc.add_paragraph("First paragraph content.")
        doc.add_paragraph("Second paragraph content.")
        doc.add_paragraph("Third paragraph content.")

        doc.add_heading("Section Two", level=2)
        doc.add_paragraph("Fourth paragraph in section two.")

        path = tmp_path / "spec.docx"
        doc.save(str(path))
        return path

    @pytest.fixture
    def pipeline(self, tmp_path: Path) -> OPPPipeline:
        resource_dir = tmp_path / "resources"
        resource_dir.mkdir()
        return OPPPipeline(resource_dir)

    def test_mock_ol_tool_translates_xliff(self, tmp_path: Path):
        xliff_input = tmp_path / "input.xlf"
        xliff_output = tmp_path / "output.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>Hello</source>
      </trans-unit>
      <trans-unit id="2">
        <source>World</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool(source_lang="en", target_lang="fr")
        result = mock_ol.translate_xliff(xliff_input, xliff_output)

        assert result is True
        assert xliff_output.exists()

        tree = ET.parse(xliff_output)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)

        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))
        assert len(units) == 2

        targets = []
        for unit in units:
            target_el = find_elem(unit, ns_uri, "target")
            assert target_el is not None
            assert "_translated" in target_el.text
            targets.append(target_el.text)

        assert "Hello_translated" in targets
        assert "World_translated" in targets

    def test_mock_ol_validates_translation(self, tmp_path: Path):
        xliff_input = tmp_path / "input.xlf"
        xliff_output = tmp_path / "output.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>Test</source>
        <target>Test_translated</target>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool(source_lang="en", target_lang="fr")
        mock_ol.translate_xliff(xliff_input, xliff_output)

        is_valid, errors = mock_ol.validate_translation(xliff_output)
        assert is_valid is True
        assert len(errors) == 0

    def test_mock_ol_validation_fails_for_missing_target(self, tmp_path: Path):
        xliff_invalid = tmp_path / "invalid.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>Test</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_invalid.write_bytes(xliff_content)

        mock_ol = MockOLTool()
        is_valid, errors = mock_ol.validate_translation(xliff_invalid)

        assert is_valid is False
        assert len(errors) > 0

    def test_opp_convert_to_xliff(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        result = pipeline.process_file(sample_docx)
        assert result.format_type == FormatType.DOCX
        assert len(result.errors) == 0

        xliff_path = tmp_path / "output.xlf"

        extraction_result = ExtractionResult(
            paragraphs=[
                ParagraphData(text="Test Document", style="Heading1", level=1),
                ParagraphData(text="First paragraph content.", style="Normal", level=None),
                ParagraphData(text="Second paragraph content.", style="Normal", level=None),
            ],
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        pipeline.generate_xliff(
            extraction_result,
            xliff_path,
            source_lang="en",
            target_lang="fr",
        )

        assert xliff_path.exists()
        assert xliff_path.stat().st_size > 0

        tree = ET.parse(xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)

        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))
        assert len(units) == 3

        sources = [find_elem(u, ns_uri, "source").text for u in units]
        assert "Test Document" in sources
        assert "First paragraph content." in sources

    def test_opp_convert_to_markdown(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        result = pipeline.process_file(sample_docx)
        assert result.format_type == FormatType.DOCX

        md_path = tmp_path / "output.md"

        extraction_result = ExtractionResult(
            paragraphs=[
                ParagraphData(text="Test Document", style="Heading1", level=1),
                ParagraphData(text="First paragraph content.", style="Normal", level=None),
            ],
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        pipeline.generate_markdown(extraction_result, md_path)

        assert md_path.exists()
        content = md_path.read_text()
        assert "Test Document" in content
        assert "First paragraph content." in content

    def test_full_pipeline_opp_to_mock_ol(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        output_dir = tmp_path / "preprocess"
        output_dir.mkdir()

        result = pipeline.process_file(sample_docx)
        assert result.format_type == FormatType.DOCX
        assert len(result.errors) == 0

        xliff_path = output_dir / "spec.xlf"

        extraction_result = ExtractionResult(
            paragraphs=[
                ParagraphData(text="Test Document", style="Heading1", level=1),
                ParagraphData(text="First paragraph content.", style="Normal", level=None),
                ParagraphData(text="Second paragraph content.", style="Normal", level=None),
                ParagraphData(text="Third paragraph content.", style="Normal", level=None),
                ParagraphData(text="Section Two", style="Heading2", level=2),
                ParagraphData(text="Fourth paragraph in section two.", style="Normal", level=None),
            ],
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        pipeline.generate_xliff(
            extraction_result,
            xliff_path,
            source_lang="en",
            target_lang="ja",
        )

        assert xliff_path.exists()

        translated_xliff_path = output_dir / "spec_ja.xlf"

        mock_ol = MockOLTool(source_lang="en", target_lang="ja")
        success = mock_ol.translate_xliff(xliff_path, translated_xliff_path)

        assert success is True
        assert translated_xliff_path.exists()

        is_valid, errors = mock_ol.validate_translation(translated_xliff_path)
        assert is_valid is True, f"Validation errors: {errors}"

        tree = ET.parse(translated_xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)

        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))
        assert len(units) == 6

        for unit in units:
            target_el = find_elem(unit, ns_uri, "target")
            assert target_el is not None
            assert "_translated" in target_el.text

    def test_md_xliff_paragraph_alignment(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        output_dir = tmp_path / "preprocess"
        output_dir.mkdir()

        paragraphs = [
            ParagraphData(text="Test Document", style="Heading1", level=1),
            ParagraphData(text="First paragraph content.", style="Normal", level=None),
            ParagraphData(text="Second paragraph content.", style="Normal", level=None),
        ]

        md_path = output_dir / "spec.md"
        xliff_path = output_dir / "spec.xlf"

        extraction_result = ExtractionResult(
            paragraphs=paragraphs,
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        pipeline.generate_markdown(extraction_result, md_path)
        pipeline.generate_xliff(extraction_result, xliff_path, source_lang="en", target_lang="fr")

        md_content = md_path.read_text()
        tree = ET.parse(xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))

        assert len(units) == len(paragraphs)

        for i, para in enumerate(paragraphs):
            source_text = find_elem(units[i], ns_uri, "source").text
            assert source_text == para.text

            expected_in_md = para.text in md_content or f"#{'#' * (para.level or 0)} {para.text}" in md_content
            assert expected_in_md, f"Paragraph '{para.text}' not found in MD"

        assert len(units) == 3

    def test_bdd_scenario_convert_both_formats(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        output_dir = tmp_path / "preprocess"
        output_dir.mkdir()

        paragraphs = [
            ParagraphData(text="Test Document", style="Heading1", level=1),
            ParagraphData(text="First paragraph content.", style="Normal", level=None),
            ParagraphData(text="Second paragraph content.", style="Normal", level=None),
            ParagraphData(text="Third paragraph content.", style="Normal", level=None),
        ]

        extraction_result = ExtractionResult(
            paragraphs=paragraphs,
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        md_path = output_dir / "spec.md"
        xliff_path = output_dir / "spec.xlf"

        pipeline.generate_markdown(extraction_result, md_path)
        pipeline.generate_xliff(extraction_result, xliff_path, source_lang="en", target_lang="ja")

        assert md_path.exists(), "spec.md should be generated"
        assert xliff_path.exists(), "spec.xlf should be generated"

        md_content = md_path.read_text()
        assert "Test Document" in md_content
        assert "First paragraph content." in md_content
        assert "Second paragraph content." in md_content
        assert "Third paragraph content." in md_content

        tree = ET.parse(xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))

        assert len(units) == 4, "Each paragraph should have a corresponding trans-unit"

        for i, para in enumerate(paragraphs):
            source_text = find_elem(units[i], ns_uri, "source").text
            assert source_text == para.text, f"XLIFF unit {i} source should match paragraph"

    def test_bdd_scenario_full_localization_pipeline(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        output_dir = tmp_path / "preprocess"
        output_dir.mkdir()

        paragraphs = [
            ParagraphData(text="Introduction", style="Heading1", level=1),
            ParagraphData(text="This is the first section.", style="Normal", level=None),
            ParagraphData(text="Details about the feature.", style="Normal", level=None),
        ]

        extraction_result = ExtractionResult(
            paragraphs=paragraphs,
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        source_xliff_path = output_dir / "spec.xlf"
        pipeline.generate_xliff(extraction_result, source_xliff_path, source_lang="en", target_lang="ja")

        translated_xliff_path = output_dir / "spec_ja.xlf"

        mock_ol = MockOLTool(source_lang="en", target_lang="ja")
        success = mock_ol.translate_xliff(source_xliff_path, translated_xliff_path)

        assert success is True
        assert translated_xliff_path.exists()

        tree = ET.parse(translated_xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)

        file_el = find_elem(root, ns_uri, "file")
        assert file_el is not None
        assert file_el.get("target-language") == "ja"

        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))
        assert len(units) == 3

        for i, para in enumerate(paragraphs):
            source_el = find_elem(units[i], ns_uri, "source")
            target_el = find_elem(units[i], ns_uri, "target")

            assert source_el is not None
            assert source_el.text == para.text
            assert target_el is not None
            assert "_translated" in target_el.text
            assert para.text in target_el.text

    def test_mock_ol_cli_interface(self, tmp_path: Path):
        from tests.mock_ol import main as mock_ol_main

        xliff_input = tmp_path / "input.xlf"
        xliff_output = tmp_path / "output.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>CLI Test</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        exit_code = mock_ol_main([str(xliff_input), str(xliff_output), "--source-lang", "en", "--target-lang", "ja"])

        assert exit_code == 0
        assert xliff_output.exists()

        tree = ET.parse(xliff_output)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        target_el = find_elem(root, ns_uri, "target")
        assert target_el is not None
        assert "CLI Test_translated" == target_el.text

    def test_xliff_content_matches_md_content(self, pipeline: OPPPipeline, sample_docx: Path, tmp_path: Path):
        output_dir = tmp_path / "output"
        output_dir.mkdir()

        paragraphs = [
            ParagraphData(text="Header One", style="Heading1", level=1),
            ParagraphData(text="Content line one.", style="Normal", level=None),
            ParagraphData(text="Content line two.", style="Normal", level=None),
        ]

        md_path = output_dir / "content.md"
        xliff_path = output_dir / "content.xlf"

        extraction_result = ExtractionResult(
            paragraphs=paragraphs,
            tables=[],
            images=[],
            metadata=DocumentMetadata(format_type="DOCX"),
        )

        pipeline.generate_markdown(extraction_result, md_path)
        pipeline.generate_xliff(extraction_result, xliff_path, source_lang="en", target_lang="de")

        md_text = md_path.read_text()
        tree = ET.parse(xliff_path)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        units = list(root.iter(f"{{{ns_uri}}}trans-unit"))

        md_paragraphs = [p.text for p in paragraphs]

        for i, unit in enumerate(units):
            source_text = find_elem(unit, ns_uri, "source").text
            assert source_text == md_paragraphs[i], f"Unit {i} source should match MD paragraph"

        assert len(units) == len(paragraphs)

    def test_mock_ol_handles_empty_xliff(self, tmp_path: Path):
        xliff_input = tmp_path / "empty.xlf"
        xliff_output = tmp_path / "empty_out.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="empty.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool()
        success = mock_ol.translate_xliff(xliff_input, xliff_output)

        assert success is True
        assert xliff_output.exists()

        is_valid, errors = mock_ol.validate_translation(xliff_output)
        assert is_valid is True

    def test_mock_ol_preserves_xml_structure(self, tmp_path: Path):
        xliff_input = tmp_path / "structured.xlf"
        xliff_output = tmp_path / "structured_out.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="document.docx" source-language="en" target-language="de" datatype="office-open-xml">
    <header>
      <prop-group>
        <prop prop-type="x-category">manual</prop>
      </prop-group>
    </header>
    <body>
      <trans-unit id="u1" translate="yes">
        <source>Important text</source>
        <note from="OPP">Context note</note>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool(source_lang="en", target_lang="de")
        success = mock_ol.translate_xliff(xliff_input, xliff_output)

        assert success is True

        tree = ET.parse(xliff_output)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)

        file_el = find_elem(root, ns_uri, "file")
        assert file_el is not None
        assert file_el.get("original") == "document.docx"
        assert file_el.get("source-language") == "en"
        assert file_el.get("target-language") == "de"
        assert file_el.get("datatype") == "office-open-xml"

        unit = find_elem(root, ns_uri, "trans-unit")
        assert unit is not None
        assert unit.get("id") == "u1"
        assert unit.get("translate") == "yes"

        note_el = find_elem(unit, ns_uri, "note")
        assert note_el is not None
        assert note_el.get("from") == "OPP"

        target_el = find_elem(unit, ns_uri, "target")
        assert target_el is not None
        assert "Important text_translated" == target_el.text


class TestOPPtoOLPipelineEdgeCases:
    def test_xliff_with_special_characters(self, tmp_path: Path):
        xliff_input = tmp_path / "special.xlf"
        xliff_output = tmp_path / "special_out.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>Tom &amp; Jerry &lt;test&gt;</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool()
        success = mock_ol.translate_xliff(xliff_input, xliff_output)

        assert success is True
        assert xliff_output.exists()

    def test_xliff_with_unicode_content(self, tmp_path: Path):
        xliff_input = tmp_path / "unicode.xlf"
        xliff_output = tmp_path / "unicode_out.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="ja" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>\xe4\xb8\xad\xe6\x96\x87\xe5\xad\x97</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        mock_ol = MockOLTool(source_lang="en", target_lang="ja")
        success = mock_ol.translate_xliff(xliff_input, xliff_output)

        assert success is True

    def test_multiple_translation_rounds(self, tmp_path: Path):
        xliff_input = tmp_path / "round1.xlf"

        xliff_content = b"""<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.txt" source-language="en" target-language="fr" datatype="plaintext">
    <body>
      <trans-unit id="1">
        <source>Original text</source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        xliff_input.write_bytes(xliff_content)

        round1_output = tmp_path / "round1_out.xlf"
        mock_ol = MockOLTool(source_lang="en", target_lang="fr")
        mock_ol.translate_xliff(xliff_input, round1_output)

        tree = ET.parse(round1_output)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        target1 = find_elem(root, ns_uri, "target").text
        assert target1 == "Original text_translated"

        round2_output = tmp_path / "round2_out.xlf"
        mock_ol.translate_xliff(round1_output, round2_output)

        tree = ET.parse(round2_output)
        root = tree.getroot()
        ns_uri = detect_xliff_ns(root)
        target2 = find_elem(root, ns_uri, "target").text
        assert target2 == target1