"""Issue #79 regression: PPTX table cells must survive into XLIFF and Markdown.

A PPTX table is a ``GraphicFrame`` with ``has_table=True`` and
``has_text_frame=False``. ``extract_shapes`` only handled GROUP shapes and
text frames, so a table matched neither branch and its cells were dropped from
BOTH the XLIFF (``table_cells`` empty) and the Markdown (``tables=[]``). These
tests drive the real ``PPTXExtractor``, ``XLIFFFileGenerator`` and
``MarkdownGenerator``.

OPP#80 Wave 2: an ``OPP_TABLE_PARAGRAPH_UNITS`` class below additionally locks
the flag-gated per-paragraph table-cell enumeration (``table_{t}_r{r}_c{c}_para{p}``
resnames) and the invariant that ``TableData``/markdown is byte-identical in
both flag states.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from opp.extractors.pptx import PPTXExtractor
from opp.markdown.generator import MarkdownGenerator
from opp.xliff import XLIFFFileGenerator

_TABLE_SEPARATOR = re.compile(r"^\|[-]+(\|[-]+)*\|$")


def _add_table(slide, data: list[list[str]]):
    rows, cols = len(data), len(data[0])
    graphic_frame = slide.shapes.add_table(
        rows, cols, Inches(1), Inches(1), Inches(4), Inches(2)
    )
    table = graphic_frame.table
    for row, values in enumerate(data):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    return graphic_frame


def _add_textbox(slide, text: str):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.5))
    box.text_frame.text = text
    return box


def _save(prs: Presentation, tmp_path: Path) -> Path:
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def _generate_xliff(result, tmp_path: Path) -> bytes:
    out_path = tmp_path / "out.xlf"
    gen = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
    gen.write_to_file(out_path)
    return out_path.read_bytes()


def _resnames(xliff_bytes: bytes) -> list[str]:
    return [r.decode("utf-8") for r in re.findall(rb'resname="([^"]+)"', xliff_bytes)]


def _unit_count(xliff_bytes: bytes) -> int:
    return xliff_bytes.count(b"<trans-unit")


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


class TestPPTXTableCells:
    def test_two_by_two_table_becomes_trans_units(self, tmp_path: Path):
        """A 2x2 table yields the four ``table_0_r*_c*`` trans-units."""
        prs = Presentation()
        _add_table(_blank_slide(prs), [["A", "B"], ["C", "D"]])
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        xliff_bytes = _generate_xliff(result, tmp_path)
        resnames = set(_resnames(xliff_bytes))
        for expected in (
            "table_0_r0_c0",
            "table_0_r0_c1",
            "table_0_r1_c0",
            "table_0_r1_c1",
        ):
            assert expected in resnames, f"missing resname {expected!r}: {resnames}"

        assert len(result.tables) == 1
        assert result.tables[0].headers == ["A", "B"]
        assert result.tables[0].rows == [["C", "D"]]

    def test_table_index_accumulates_across_slides(self, tmp_path: Path):
        """The ``t`` counter is deck-wide: two slides -> table_0_* and table_1_*."""
        prs = Presentation()
        _add_table(_blank_slide(prs), [["A1", "B1"], ["C1", "D1"]])
        _add_table(_blank_slide(prs), [["A2", "B2"], ["C2", "D2"]])
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        xliff_bytes = _generate_xliff(result, tmp_path)
        resnames = set(_resnames(xliff_bytes))
        assert "table_0_r0_c0" in resnames
        assert "table_1_r0_c0" in resnames
        assert not any(name.startswith("table_2") for name in resnames), resnames

    def test_empty_cell_produces_no_unit(self, tmp_path: Path):
        """An empty cell is skipped; only the filled cell becomes a unit."""
        prs = Presentation()
        _add_table(_blank_slide(prs), [["Filled"], [""]])
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        xliff_bytes = _generate_xliff(result, tmp_path)
        assert _unit_count(xliff_bytes) == 1
        resnames = set(_resnames(xliff_bytes))
        assert "table_0_r0_c0" in resnames
        assert "table_0_r1_c0" not in resnames

    def test_three_column_table_renders_gfm(self, tmp_path: Path):
        """``TableData`` (first row = headers) renders as a real GFM table."""
        prs = Presentation()
        _add_table(_blank_slide(prs), [["H1", "H2", "H3"], ["a", "b", "c"]])
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        markdown = MarkdownGenerator().generate(result)
        table_lines = [
            line for line in markdown.splitlines() if line.strip().startswith("|")
        ]
        assert table_lines[0] == "| H1 | H2 | H3 |"
        assert _TABLE_SEPARATOR.match(table_lines[1]), table_lines[1]
        assert table_lines[2] == "| a | b | c |"

    def test_interleaved_shape_order_is_preserved(self, tmp_path: Path):
        """paragraph -> table -> paragraph must render in that exact order."""
        prs = Presentation()
        slide = _blank_slide(prs)
        _add_textbox(slide, "PRE TEXT")
        _add_table(slide, [["H1", "H2"], ["v1", "v2"]])
        _add_textbox(slide, "POST TEXT")
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        markdown = MarkdownGenerator().generate(result)
        assert (
            markdown.index("PRE TEXT")
            < markdown.index("| H1")
            < markdown.index("POST TEXT")
        ), markdown

    def test_merged_origin_cell_renders_as_br(self, tmp_path: Path):
        """A merged origin's multi-paragraph text stays inside a valid GFM table."""
        prs = Presentation()
        slide = _blank_slide(prs)
        table = _add_table(slide, [["A", "B"], ["C", "D"]]).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "A\nB"
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        assert result.table_cells[0].text == "A\nB"
        markdown = MarkdownGenerator().generate(result)
        assert "A<br>B" in markdown
        table_lines = [
            line for line in markdown.splitlines() if line.strip().startswith("|")
        ]
        assert _TABLE_SEPARATOR.match(table_lines[1]), table_lines
        assert table_lines[0].count("|") == table_lines[2].count("|") == 3

    def test_grouped_table_is_skipped(self, tmp_path: Path):
        """A table nested in a GroupShape is skipped, matching ORF's writer."""
        prs = Presentation()
        slide = _blank_slide(prs)
        group = slide.shapes.add_group_shape()
        graphic_frame = _add_table(slide, [["A", "B"], ["C", "D"]])
        # python-pptx's GroupShapes has no add_table; move the graphicFrame into
        # the group so it is no longer a top-level p:spTree child.
        group.shapes._grpSp.append(graphic_frame._element)
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        assert result.tables == []
        assert result.table_cells == []
        xliff_bytes = _generate_xliff(result, tmp_path)
        assert not any(
            name.startswith("table_") for name in _resnames(xliff_bytes)
        )


def _fill_cell_paragraphs(table, row: int, col: int, texts: list[str]):
    """Set a cell's RAW paragraphs, empty strings included, exactly as given.

    ``text_frame.text = "\\n".join(...)`` would also do this, but the explicit
    loop keeps an all-empty or empty-middle cell unambiguous regardless of
    python-pptx's newline-compression behavior.
    """
    text_frame = table.cell(row, col).text_frame
    text_frame.text = texts[0]
    for extra in texts[1:]:
        p = text_frame.add_paragraph()
        if extra:
            p.text = extra


def _cell_summary(result) -> list[tuple[int, int, int, int | None, str]]:
    return [
        (c.table_index, c.row, c.col, c.para_index, c.text)
        for c in result.table_cells
    ]


class TestPPTXTableParagraphUnits:
    """OPP#80 Wave 2: ``OPP_TABLE_PARAGRAPH_UNITS``-gated per-paragraph units.

    The flag is a helper function in ``opp.config`` that reads the env var at
    CALL time, so ``monkeypatch.setenv``/``delenv`` is the contract-supported
    toggle. Every test re-extracts from a fresh deck — there is no import-time
    caching to invalidate.
    """

    def _deck_with_cell(
        self, tmp_path: Path, texts: list[str]
    ) -> object:
        prs = Presentation()
        slide = _blank_slide(prs)
        graphic_frame = _add_table(slide, [["X"]])
        _fill_cell_paragraphs(graphic_frame.table, 0, 0, texts)
        path = _save(prs, tmp_path)
        return PPTXExtractor().extract(path)

    def test_flag_off_three_paragraph_cell_is_one_whole_unit(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag OFF (default): legacy whole-cell unit, bare resname, no shift.

        Pins BOTH the legacy XLIFF resname AND the unchanged markdown ``TableData``
        for a multi-paragraph cell.
        """
        monkeypatch.delenv("OPP_TABLE_PARAGRAPH_UNITS", raising=False)
        result = self._deck_with_cell(tmp_path, ["P0", "P1", "P2"])

        assert _cell_summary(result) == [(0, 0, 0, None, "P0\nP1\nP2")]
        assert result.tables[0].headers == ["P0\nP1\nP2"]
        assert result.tables[0].rows == []
        markdown = MarkdownGenerator().generate(result)
        pipe_lines = [
            line for line in markdown.splitlines() if line.strip().startswith("|")
        ]
        assert pipe_lines[0] == "| P0<br>P1<br>P2 |"
        xliff_bytes = _generate_xliff(result, tmp_path)
        assert _resnames(xliff_bytes) == ["table_0_r0_c0"]

    def test_flag_on_three_paragraph_cell_emits_para_units(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON: one ``_para{p}`` unit per raw paragraph, including 0.

        The markdown assertions are the anti-regression guard for the trap in
        OPP#80: the per-paragraph units feed ONLY ``table_cells`` — ``_cell_text``
        and the ``all_rows``/``TableData`` construction must stay untouched, so
        the markdown table still has exactly ONE cell entry with the joined text.
        """
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        result = self._deck_with_cell(tmp_path, ["P0", "P1", "P2"])

        assert _cell_summary(result) == [
            (0, 0, 0, 0, "P0"),
            (0, 0, 0, 1, "P1"),
            (0, 0, 0, 2, "P2"),
        ]
        assert result.tables[0].headers == ["P0\nP1\nP2"]
        assert result.tables[0].rows == []
        markdown = MarkdownGenerator().generate(result)
        pipe_lines = [
            line for line in markdown.splitlines() if line.strip().startswith("|")
        ]
        assert pipe_lines[0] == "| P0<br>P1<br>P2 |"

    def test_flag_on_single_paragraph_cell_is_bare(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON, single-paragraph cell: one bare whole-cell unit, no suffix."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        result = self._deck_with_cell(tmp_path, ["ONLY"])

        assert _cell_summary(result) == [(0, 0, 0, None, "ONLY")]

    def test_flag_on_empty_paragraph_consumes_index(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON, cell [P0, '', P2]: indices 0 and 2; the empty p=1 gets NO unit.

        ``para_index`` counts the RAW paragraph list INCLUDING empty paragraphs —
        skipping the empty one inline would shift every later index and make ORF
        write the wrong paragraph.
        """
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        result = self._deck_with_cell(tmp_path, ["P0", "", "P2"])

        assert _cell_summary(result) == [
            (0, 0, 0, 0, "P0"),
            (0, 0, 0, 2, "P2"),
        ]

    def test_flag_on_all_empty_cell_emits_no_unit(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON, all-empty cell: no unit (same as flag OFF)."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        result = self._deck_with_cell(tmp_path, ["", "", ""])

        assert result.table_cells == []

    def test_flag_on_resnames_end_to_end_through_generator(
        self, tmp_path: Path, monkeypatch
    ):
        """Real generator: ``table_0_r0_c0_para0/_para1/_para2`` ON, bare OFF."""
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        result = self._deck_with_cell(tmp_path, ["P0", "P1", "P2"])
        xliff_bytes = _generate_xliff(result, tmp_path)
        assert set(_resnames(xliff_bytes)) == {
            "table_0_r0_c0_para0",
            "table_0_r0_c0_para1",
            "table_0_r0_c0_para2",
        }

        monkeypatch.delenv("OPP_TABLE_PARAGRAPH_UNITS", raising=False)
        result_off = self._deck_with_cell(tmp_path, ["P0", "P1", "P2"])
        xliff_off = _generate_xliff(result_off, tmp_path)
        assert set(_resnames(xliff_off)) == {"table_0_r0_c0"}

    def test_flag_on_merged_cell_follows_raw_indices(
        self, tmp_path: Path, monkeypatch
    ):
        """Flag ON + merged cells: no crash, indices follow the raw ``a:p`` list.

        The origin ``a:tc`` holds paragraphs 'A', 'B' at raw indices 0/1; the
        spanned ``a:tc`` (raw col 1, empty) contributes no unit. Raw coordinates
        come from ``findall(a:tr)``/``findall(a:tc)`` matching ORF's writer.
        """
        monkeypatch.setenv("OPP_TABLE_PARAGRAPH_UNITS", "1")
        prs = Presentation()
        slide = _blank_slide(prs)
        table = _add_table(slide, [["A", "B"], ["C", "D"]]).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "A\nB"
        result = PPTXExtractor().extract(_save(prs, tmp_path))

        assert _cell_summary(result) == [
            (0, 0, 0, 0, "A"),
            (0, 0, 0, 1, "B"),
            (0, 1, 0, None, "C"),
            (0, 1, 1, None, "D"),
        ]
        xliff_bytes = _generate_xliff(result, tmp_path)
        assert "table_0_r0_c0_para0".encode("utf-8") in xliff_bytes
        assert "table_0_r0_c0_para1".encode("utf-8") in xliff_bytes
