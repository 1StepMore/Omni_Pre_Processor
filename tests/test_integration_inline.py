import tempfile
from pathlib import Path
from docx import Document
from pptx import Presentation
from opp.extractors.docx import DOCXExtractor
from opp.extractors.pptx import PPTXExtractor
from opp.extractors.epub import EPUBExtractor
from opp.extractors.html import HTMLExtractor
from opp.xliff.generator import XLIFFFileGenerator
from opp.utils.dataclasses import RunData


class TestIntegrationInline:
    def test_docx_bold_roundtrip(self):
        doc = Document()
        para = doc.add_paragraph()
        run1 = para.add_run("Hello ")
        run1.bold = True
        run2 = para.add_run("World")
        run2.bold = False

        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = Path(tmpdir) / "test.docx"
            doc.save(str(docx_path))

            extractor = DOCXExtractor()
            result = extractor.extract(docx_path)

            generator = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
            xliff_bytes = generator.to_bytes()
            xliff_str = xliff_bytes.decode('utf-8')

            assert '<bx id="1" type="bold"/>' in xliff_str
            assert '&lt;bx' not in xliff_str

    def test_pptx_bold_roundtrip(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        txBox = slide.shapes.add_textbox(100, 100, 300, 50)
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = "Hello "
        r1.font.bold = True
        r2 = p.add_run()
        r2.text = "World"
        r2.font.bold = False

        with tempfile.TemporaryDirectory() as tmpdir:
            pptx_path = Path(tmpdir) / "test.pptx"
            prs.save(str(pptx_path))

            extractor = PPTXExtractor()
            result = extractor.extract(pptx_path)

            generator = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
            xliff_bytes = generator.to_bytes()
            xliff_str = xliff_bytes.decode('utf-8')

            assert '<bx id="1" type="bold"/>' in xliff_str
            assert '&lt;bx' not in xliff_str

    def test_epub_italic_roundtrip(self):
        from bs4 import BeautifulSoup
        from opp.utils.dataclasses import ParagraphData

        html_content = '<p><em>Italic Text</em></p>'
        soup = BeautifulSoup(html_content, 'html.parser')
        element = soup.find('em')

        extractor = EPUBExtractor()
        runs = extractor.extract_runs(element)

        para = ParagraphData(text="Italic Text", runs=runs)

        from opp.utils.dataclasses import ExtractionResult
        from opp.utils.dataclasses import DocumentMetadata, ImageData

        result = ExtractionResult(
            paragraphs=[para],
            tables=[],
            images=[],
            metadata=DocumentMetadata(),
        )

        generator = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
        xliff_bytes = generator.to_bytes()
        xliff_str = xliff_bytes.decode('utf-8')

        assert '<bx id="1" type="italic"/>' in xliff_str
        assert '&lt;bx' not in xliff_str

    def test_html_underline_roundtrip(self):
        from bs4 import BeautifulSoup
        from opp.utils.dataclasses import ParagraphData

        html_content = '<p><u>Underlined</u></p>'
        soup = BeautifulSoup(html_content, 'html.parser')
        element = soup.find('u')

        extractor = HTMLExtractor()
        runs = extractor.extract_runs(element)

        para = ParagraphData(text="Underlined", runs=runs)

        from opp.utils.dataclasses import ExtractionResult
        from opp.utils.dataclasses import DocumentMetadata, ImageData

        result = ExtractionResult(
            paragraphs=[para],
            tables=[],
            images=[],
            metadata=DocumentMetadata(),
        )

        generator = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
        xliff_bytes = generator.to_bytes()
        xliff_str = xliff_bytes.decode('utf-8')

        assert '<bx id="1" type="underline"/>' in xliff_str
        assert '&lt;bx' not in xliff_str

    def test_no_inline_elements_plain_text(self):
        from opp.utils.dataclasses import ParagraphData, ExtractionResult, DocumentMetadata, ImageData

        para = ParagraphData(text="Plain text without formatting")

        result = ExtractionResult(
            paragraphs=[para],
            tables=[],
            images=[],
            metadata=DocumentMetadata(),
        )

        generator = XLIFFFileGenerator.from_extraction_result(result, "en", "zh")
        xliff_bytes = generator.to_bytes()
        xliff_str = xliff_bytes.decode('utf-8')

        assert '<bx' not in xliff_str
        assert '<ex' not in xliff_str
        assert 'Plain text without formatting' in xliff_str