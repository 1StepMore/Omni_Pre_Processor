from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import PP_PLACEHOLDER, MSO_SHAPE_TYPE

from opp.extractors.pptx import PPTXExtractor


def create_pptx_with_title_shape(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    title_shape = slide.shapes.title
    title_shape.text = "Slide Title Text"

    body_shape = slide.shapes.placeholders[1]
    body_shape.text = "Body text content"

    pptx_path = tmp_path / "title_test.pptx"
    prs.save(str(pptx_path))
    return pptx_path


def create_pptx_with_body_only(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    body_shape = slide.shapes.placeholders[1]
    body_shape.text = "Body only content"

    pptx_path = tmp_path / "body_only.pptx"
    prs.save(str(pptx_path))
    return pptx_path


class TestPPTXTitleShapeMapping:
    def test_title_shape_gets_level_one(self, tmp_path: Path):
        pptx_path = create_pptx_with_title_shape(tmp_path)

        extractor = PPTXExtractor()
        result = extractor.extract(pptx_path)

        title_paragraph = None
        for p in result.paragraphs:
            if p.text == "Slide Title Text":
                title_paragraph = p
                break

        assert title_paragraph is not None, "Should find title paragraph"
        assert title_paragraph.level == 1, f"TITLE shape should have level=1, got {title_paragraph.level}"
        assert title_paragraph.style is not None and "Heading" in title_paragraph.style, \
            f"TITLE shape style should contain 'Heading', got {title_paragraph.style}"

    def test_body_shape_keeps_level_none(self, tmp_path: Path):
        pptx_path = create_pptx_with_body_only(tmp_path)

        extractor = PPTXExtractor()
        result = extractor.extract(pptx_path)

        body_paragraph = None
        for p in result.paragraphs:
            if p.text == "Body only content":
                body_paragraph = p
                break

        assert body_paragraph is not None, "Should find body paragraph"
        assert body_paragraph.level is None, f"BODY shape should have level=None, got {body_paragraph.level}"

    def test_group_shape_text_gets_level_none(self, sample_files_edge: Path):
        pptx_path = sample_files_edge / "grouped.pptx"

        extractor = PPTXExtractor()
        result = extractor.extract(pptx_path)

        texts = [p.text for p in result.paragraphs]
        assert "Shape 1" in texts or "Shape 2" in texts, "Should find grouped shapes"

        for p in result.paragraphs:
            if p.text in ("Shape 1", "Shape 2"):
                assert p.level is None, f"GROUP shape text should have level=None, got {p.level}"
                break