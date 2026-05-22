"""Unit tests for DOCX run extraction (inline formatting)."""

import pytest
from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from opp.extractors.docx import DOCXExtractor
from opp.utils.dataclasses import RunData


class TestDOCXExtractRuns:
    """Tests for DOCXExtractor.extract_runs() method."""

    def _create_para_with_runs(self, runs_data):
        """Helper to create a paragraph with runs from tuples of (text, bold, italic, underline, strike)."""
        doc = Document()
        para = doc.add_paragraph()
        for text, bold, italic, underline, strike in runs_data:
            run = para.add_run(text)
            run.bold = bold
            run.italic = italic
            run.underline = underline
            run.strikethrough = strike
        return para

    def test_extract_runs_single_plain(self):
        """Single run with no formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Hello World")
        # No formatting set

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Hello World"
        assert runs[0].bold is False
        assert runs[0].italic is False
        assert runs[0].underline is False
        assert runs[0].strike is False

    def test_extract_runs_single_bold(self):
        """Single run with bold formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Bold Text")
        run.bold = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Bold Text"
        assert runs[0].bold is True
        assert runs[0].italic is False
        assert runs[0].underline is False
        assert runs[0].strike is False

    def test_extract_runs_single_italic(self):
        """Single run with italic formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Italic Text")
        run.italic = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Italic Text"
        assert runs[0].bold is False
        assert runs[0].italic is True

    def test_extract_runs_single_underline(self):
        """Single run with underline formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Underlined Text")
        run.underline = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Underlined Text"
        assert runs[0].underline is True

    def test_extract_runs_single_strike(self):
        """Single run with strikethrough formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Strikethrough Text")
        run.font.strike = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Strikethrough Text"
        assert runs[0].strike is True

    def test_extract_runs_combined_formatting(self):
        """Single run with multiple formatting (bold + italic)."""
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Bold Italic")
        run.bold = True
        run.italic = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Bold Italic"
        assert runs[0].bold is True
        assert runs[0].italic is True

    def test_extract_runs_multiple_plain(self):
        """Multiple runs without formatting - whitespace runs are skipped."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("Hello")
        para.add_run(" ")
        para.add_run("World")

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 2
        assert runs[0].text == "Hello"
        assert runs[1].text == "World"

    def test_extract_runs_multiple_formatted(self):
        """Multiple runs with different formatting."""
        doc = Document()
        para = doc.add_paragraph()
        run1 = para.add_run("Hello ")
        run1.bold = True
        run2 = para.add_run("World")
        run2.italic = True

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 2
        assert runs[0].text == "Hello "
        assert runs[0].bold is True
        assert runs[1].text == "World"
        assert runs[1].italic is True

    def test_extract_runs_empty_text(self):
        """Run with empty text should be skipped."""
        doc = Document()
        para = doc.add_paragraph()
        run1 = para.add_run("")
        run2 = para.add_run("Actual Text")
        run3 = para.add_run("")

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Actual Text"

    def test_extract_runs_whitespace_only(self):
        """Run with whitespace only should be skipped."""
        doc = Document()
        para = doc.add_paragraph()
        run1 = para.add_run("   ")
        run2 = para.add_run("Text")

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 1
        assert runs[0].text == "Text"

    def test_extract_runs_no_runs(self):
        """Paragraph with no runs."""
        doc = Document()
        para = doc.add_paragraph()
        # Empty paragraph has no runs

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 0

    def test_extract_runs_para_with_only_empty_runs(self):
        """Paragraph where all runs are empty."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("")
        para.add_run("   ")

        extractor = DOCXExtractor()
        runs = extractor.extract_runs(para)

        assert len(runs) == 0