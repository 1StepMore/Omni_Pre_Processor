from datetime import date, datetime
from pathlib import Path

import openpyxl

from opp.extractors.base import ExtractorBase
from opp.logger import logger
from opp.utils.dataclasses import (
    ExtractionResult,
    ParagraphData,
    TableData,
)
from opp.utils.exceptions import ValidationError


class XLSXExtractor(ExtractorBase):
    MAX_ROWS_WARNING = 10000

    def supported_extensions(self) -> list[str]:
        return [".xlsx"]

    def extract_tables(self, input_path: Path) -> list[TableData]:
        self.validate_file(input_path)

        try:
            wb = openpyxl.load_workbook(input_path, data_only=True)
        except Exception as e:
            logger.debug("Failed to open XLSX in extract_tables: %s", e)
            if "password" in str(e).lower():
                raise ValidationError(f"文件受密码保护: {input_path}")
            raise ValueError(f"无法打开XLSX文件: {input_path}")

        result: list[TableData] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers: list[str] = []
            rows: list[list[str]] = []

            max_row = ws.max_row or 0
            max_col = ws.max_column or 0

            if max_row == 0 or max_col == 0:
                continue

            # Extract headers from first row
            header_row_has_data = False
            for cell in ws[1]:
                val = cell.value
                if val is None:
                    headers.append("")
                elif isinstance(val, (datetime, date)):
                    headers.append(val.isoformat())
                    header_row_has_data = True
                else:
                    headers.append(str(val))
                    header_row_has_data = True

            # Skip sheet if header row has no data
            if not header_row_has_data:
                continue

            # Extract data rows (skip header row)
            for row_idx in range(2, max_row + 1):
                row_values: list[str] = []
                for col_idx in range(1, max_col + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    val = cell.value
                    if val is None:
                        row_values.append("")
                    elif isinstance(val, (datetime, date)):
                        row_values.append(val.isoformat())
                    else:
                        row_values.append(str(val))
                rows.append(row_values)

            if headers or rows:
                result.append(TableData(headers=headers, rows=rows))

        return result

    def extract(self, input_path: Path) -> ExtractionResult:
        self.validate_file(input_path)
        metadata = self.get_file_info(input_path)
        warnings: list[str] = []

        try:
            wb = openpyxl.load_workbook(input_path, data_only=True)
        except Exception as e:
            logger.debug("Failed to open XLSX in extract: %s", e)
            if "password" in str(e).lower():
                raise ValidationError(f"文件受密码保护: {input_path}")
            raise ValueError(f"无法打开XLSX文件: {input_path}")

        paragraphs: list[ParagraphData] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            paragraphs.append(ParagraphData(
                text=f"=== Sheet: {sheet_name} ===",
                style="sheet_header",
            ))

            if ws.max_row > self.MAX_ROWS_WARNING:
                warnings.append(f"工作表 '{sheet_name}' 超过 {self.MAX_ROWS_WARNING} 行")

            max_col = ws.max_column or 1
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col):
                row_values = []
                for cell in row:
                    val = cell.value
                    if val is None:
                        row_values.append("")
                    elif isinstance(val, (datetime, date)):
                        row_values.append(val.isoformat())
                    else:
                        row_values.append(str(val))
                line = " | ".join(row_values).strip()
                if line:
                    paragraphs.append(ParagraphData(
                        text=line,
                        style=None,
                    ))

        if not any(p.style != "sheet_header" for p in paragraphs):
            warnings.append("工作簿为空")

        return ExtractionResult(
            paragraphs=paragraphs,
            tables=[],
            images=[],
            metadata=metadata,
            warnings=warnings,
        )