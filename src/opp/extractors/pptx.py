from dataclasses import replace
from pathlib import Path
import zipfile
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from opp.extractors.base import ExtractorBase
from opp.utils.dataclasses import (
    ExtractionResult,
    ImageData,
    ParagraphData,
    RunData,
    SlideData,
)
from opp.utils.exceptions import CorruptedFileError, UnsupportedFormatError


class PPTXExtractor(ExtractorBase):
    def supported_extensions(self) -> list[str]:
        return [".pptx"]

    def extract(self, input_path: Path) -> ExtractionResult:
        self.validate_file(input_path)
        metadata = self.get_file_info(input_path)
        warnings: list[str] = []

        try:
            prs = Presentation(str(input_path))
        except Exception:
            if ".pptm" in str(input_path).lower():
                raise UnsupportedFormatError(f"不支持的PPTX格式（宏已启用）: {input_path}")
            raise CorruptedFileError(f"文件损坏或无法解析: {input_path}")

        slides = self.extract_slides(prs)
        images = self.extract_images(prs)

        skeleton_bytes: bytes | None = None
        skeleton_files: list[str] | None = None
        try:
            with open(input_path, 'rb') as f:
                skeleton_bytes = f.read()
            with zipfile.ZipFile(input_path, 'r') as zf:
                skeleton_files = [f for f in zf.namelist() if f.startswith('ppt/')]
        except zipfile.BadZipFile:
            warnings.append("Skeleton extraction failed: not a valid ZIP/PPTX file")
            skeleton_bytes = None
            skeleton_files = None

        all_paragraphs: list[ParagraphData] = []
        for slide_data in slides:
            all_paragraphs.extend(slide_data.shapes)

        if not all_paragraphs:
            warnings.append("演示文稿为空")

        metadata = replace(metadata, page_count=len(prs.slides))

        return ExtractionResult(
            paragraphs=all_paragraphs,
            tables=[],
            images=images,
            metadata=metadata,
            warnings=warnings,
            skeleton=skeleton_bytes,
            skeleton_files=skeleton_files,
        )

    def extract_slides(self, prs: Presentation) -> list[SlideData]:
        result: list[SlideData] = []
        for i, slide in enumerate(prs.slides):
            shapes = self.extract_shapes(slide)
            notes = self.extract_notes(slide)
            result.append(SlideData(
                index=i,
                shapes=shapes,
                notes=notes,
            ))
        return result

    def _is_title_shape(self, shape: Any) -> bool:
        if not shape.is_placeholder:
            return False
        try:
            return shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        except (ValueError, AttributeError):
            return False

    def extract_runs(self, shape: Any) -> list[RunData]:
        runs = []
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                text = run.text
                if not text or not text.strip():
                    continue
                font = run.font
                font_size_hp = None
                if font.size is not None:
                    font_size_hp = int(font.size.pt * 2)
                color_hex = None
                if font.color.type is not None and font.color.rgb is not None:
                    color_hex = str(font.color.rgb)
                run_data = RunData(
                    text=text,
                    bold=bool(font.bold) if font.bold else False,
                    italic=bool(font.italic) if font.italic else False,
                    underline=bool(font.underline) if font.underline else False,
                    strike=False,
                    font_size=font_size_hp,
                    font_name=font.name,
                    color=color_hex,
                )
                runs.append(run_data)
        return runs

    def extract_shapes(self, slide: Any) -> list[ParagraphData]:
        result: list[ParagraphData] = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                result.extend(self._flatten_group(shape))
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    runs = self.extract_runs(shape)
                    plain_text = ''.join(r.text for r in runs)
                    if self._is_title_shape(shape):
                        result.append(ParagraphData(
                            text=plain_text,
                            style="Heading 1",
                            level=1,
                            runs=runs,
                        ))
                    else:
                        result.append(ParagraphData(
                            text=plain_text,
                            style=shape.shape_type.name if hasattr(shape.shape_type, 'name') else None,
                            level=None,
                            runs=runs,
                        ))
        return result

    def _flatten_group(self, group: Any) -> list[ParagraphData]:
        result: list[ParagraphData] = []
        for shape in group.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                result.extend(self._flatten_group(shape))
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    runs = self.extract_runs(shape)
                    plain_text = ''.join(r.text for r in runs)
                    if self._is_title_shape(shape):
                        result.append(ParagraphData(
                            text=plain_text,
                            style="Heading 1",
                            level=1,
                            runs=runs,
                        ))
                    else:
                        result.append(ParagraphData(
                            text=plain_text,
                            style=shape.shape_type.name if hasattr(shape.shape_type, 'name') else None,
                            level=None,
                            runs=runs,
                        ))
        return result

    def extract_notes(self, slide: Any) -> str:
        notes_slide = slide.notes_slide
        if notes_slide and notes_slide.notes_text_frame:
            return notes_slide.notes_text_frame.text.strip()
        return ""

    def extract_images(self, prs: Presentation) -> list[ImageData]:
        result: list[ImageData] = []
        for slide_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if hasattr(shape, "image"):
                    try:
                        image = shape.image
                        result.append(ImageData(
                            data=image.blob,
                            mime_type=image.content_type,
                            width=image.size.width if hasattr(image.size, 'width') else None,
                            height=image.size.height if hasattr(image.size, 'height') else None,
                            slide_index=slide_idx,
                        ))
                    except Exception as e:
                        logger.warning(f"Failed to extract image from PPTX: {e}")
        return result
